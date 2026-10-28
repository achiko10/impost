import os
from functools import wraps
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.template.loader import render_to_string
from django.utils import timezone
from datetime import datetime
from io import BytesIO
from rest_framework_simplejwt.authentication import JWTAuthentication
from rest_framework_simplejwt.tokens import AccessToken

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from apps.tasks.models import Task
from apps.tasks.serializers import TaskSerializer
from apps.companies.models import Company, Equipment
from apps.companies.serializers import EquipmentSerializer
from apps.accounts.models import User, get_user_tenant


def report_auth_required(view_func):
    """Allow Django Session auth, Authorization: Bearer <jwt>, or ?token=<jwt> for reports."""

    @wraps(view_func)
    def _wrapped(request, *args, **kwargs):
        if request.user and request.user.is_authenticated:
            return view_func(request, *args, **kwargs)

        token_str = request.GET.get("token")
        if not token_str:
            auth_header = request.META.get("HTTP_AUTHORIZATION", "")
            if auth_header.startswith("Bearer "):
                token_str = auth_header.split(" ", 1)[1].strip()

        if token_str:
            try:
                validated = AccessToken(token_str)
                user_id = validated.get("user_id")
                user = User.objects.filter(id=user_id, is_active=True).first()
                if user:
                    request.user = user
                    return view_func(request, *args, **kwargs)
            except Exception:
                pass

        return login_required(view_func)(request, *args, **kwargs)

    return _wrapped

# Try to import xhtml2pdf for PDF generation
try:
    from xhtml2pdf import pisa
    from xhtml2pdf.default import DEFAULT_FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False


def _ensure_georgian_pdf_font():
    """Register a Unicode/Georgian TrueType font directly in ReportLab to avoid xhtml2pdf temp-file lock on Windows."""
    if not PDF_AVAILABLE:
        return
    try:
        candidates = [
            "C:/Windows/Fonts/sylfaen.ttf",
            "C:/Windows/Fonts/arial.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        ]
        for font_path in candidates:
            if os.path.exists(font_path):
                pdfmetrics.registerFont(TTFont("GeorgianFont", font_path))
                DEFAULT_FONT["helvetica"] = "GeorgianFont"
                DEFAULT_FONT["georgianfont"] = "GeorgianFont"
                DEFAULT_FONT["sans-serif"] = "GeorgianFont"
                DEFAULT_FONT["arial"] = "GeorgianFont"
                break
    except Exception:
        pass


@report_auth_required
def generate_pdf_report(request, company_id):
    """Generate PDF/printable report for a company"""
    company = get_object_or_404(Company, id=company_id)

    # Security check
    if request.user.role == "client":
        if request.user.company != company:
            return HttpResponse("Unauthorized", status=403)
    elif request.user.role == "manager":
        user_tenant = get_user_tenant(request.user)
        if company.tenant != user_tenant:
            return HttpResponse("Unauthorized", status=403)
    else:
        return HttpResponse("Unauthorized", status=403)

    tasks = (
        Task.objects.filter(schedule__equipment__site__company=company)
        .select_related("schedule__equipment__site", "assigned_to", "report")
        .order_by("-scheduled_date")
    )

    stats = {
        "total": tasks.count(),
        "pending": tasks.filter(status="pending").count(),
        "in_progress": tasks.filter(status="in_progress").count(),
        "completed": tasks.filter(status="completed").count(),
        "approved": tasks.filter(status="approved").count(),
    }

    is_pdf_download = request.GET.get("download") == "pdf"
    context = {
        "company": company,
        "tasks": tasks,
        "stats": stats,
        "generated_at": timezone.now(),
        "pdf_mode": is_pdf_download,
    }

    if is_pdf_download:
        if not PDF_AVAILABLE:
            return HttpResponse(
                "PDF generation not available. Install xhtml2pdf.", status=503
            )
        _ensure_georgian_pdf_font()
        html = render_to_string("reports/pdf_report.html", context)

        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = (
            f'attachment; filename="company_{company.id}_report_{datetime.now().strftime("%Y%m%d")}.pdf"'
        )

        pisa_status = pisa.CreatePDF(BytesIO(html.encode("utf-8")), dest=response)

        if pisa_status.err:
            return HttpResponse("PDF generation error", status=500)

        return response

    return render(request, "reports/pdf_report.html", context)


@report_auth_required
def monthly_report(request):
    """Generate monthly report"""
    if request.user.role not in ["manager"]:
        return HttpResponse("Unauthorized", status=403)

    user_tenant = get_user_tenant(request.user)
    today = timezone.now()
    start_of_month = today.replace(day=1, hour=0, minute=0, second=0, microsecond=0)

    tasks = Task.objects.filter(
        scheduled_date__gte=start_of_month,
        schedule__equipment__site__company__tenant=user_tenant,
    ).select_related("schedule__equipment__site__company", "assigned_to")

    companies_data = {}
    for task in tasks:
        company_name = task.schedule.equipment.site.company.name
        if company_name not in companies_data:
            companies_data[company_name] = {
                "total": 0,
                "completed": 0,
                "pending": 0,
            }
        companies_data[company_name]["total"] += 1
        if task.status in ["completed", "approved"]:
            companies_data[company_name]["completed"] += 1
        else:
            companies_data[company_name]["pending"] += 1

    if request.GET.get("format") == "json":
        return JsonResponse(
            {
                "month": today.strftime("%B %Y"),
                "companies_data": companies_data,
                "total_tasks": tasks.count(),
                "generated_at": today.isoformat(),
            }
        )

    context = {
        "month": today.strftime("%B %Y"),
        "companies_data": companies_data,
        "total_tasks": tasks.count(),
        "generated_at": today,
    }

    return render(request, "reports/monthly_report.html", context)


@report_auth_required
def equipment_report(request, equipment_id):
    """Generate equipment maintenance history report"""
    equipment = get_object_or_404(Equipment, id=equipment_id)

    if request.user.role == "client":
        if request.user.company != equipment.site.company:
            return HttpResponse("Unauthorized", status=403)
    elif request.user.role == "manager":
        user_tenant = get_user_tenant(request.user)
        if equipment.site.company.tenant != user_tenant:
            return HttpResponse("Unauthorized", status=403)
    elif request.user.role == "inspector":
        user_tenant = get_user_tenant(request.user)
        if equipment.site.company.tenant != user_tenant:
            return HttpResponse("Unauthorized", status=403)

    tasks = (
        Task.objects.filter(schedule__equipment=equipment)
        .select_related("assigned_to", "report", "schedule")
        .order_by("-scheduled_date")
    )

    if request.GET.get("format") == "json":
        return JsonResponse(
            {
                "equipment": EquipmentSerializer(equipment).data,
                "tasks": TaskSerializer(tasks, many=True).data,
                "generated_at": timezone.now().isoformat(),
            }
        )

    base_template = (
        "manager/base.html"
        if getattr(request.user, "role", None) == "manager"
        else "base/base.html"
    )
    context = {
        "equipment": equipment,
        "tasks": tasks,
        "generated_at": timezone.now(),
        "base_template": base_template,
    }

    return render(request, "reports/equipment_report.html", context)


@report_auth_required
def task_completion_pdf(request, task_id):
    """დასრულებული ტასკის PDF - ოფიციალური დოკუმენტი"""
    task = get_object_or_404(Task, id=task_id)

    equipment = task.schedule.equipment
    site = equipment.site
    company = site.company

    # Check permission
    if request.user.role == "inspector" and task.assigned_to != request.user:
        return HttpResponse("Unauthorized", status=403)
    elif request.user.role == "client" and request.user.company != company:
        return HttpResponse("Unauthorized", status=403)
    elif request.user.role == "manager":
        user_tenant = get_user_tenant(request.user)
        if company.tenant != user_tenant:
            return HttpResponse("Unauthorized", status=403)

    # Task must be completed
    if task.status not in ["completed", "approved"]:
        return HttpResponse("Task not completed yet", status=400)

    context = {
        "task": task,
        "equipment": equipment,
        "site": site,
        "company": company,
        "report": getattr(task, "report", None),
        "photos": (
            task.report.photos.all()
            if hasattr(task, "report") and task.report
            else []
        ),
        "generated_at": timezone.now(),
    }

    # Check if PDF download requested
    if request.GET.get("download") == "pdf" and PDF_AVAILABLE:
        _ensure_georgian_pdf_font()
        html = render_to_string("reports/task_completion_pdf.html", context)

        response = HttpResponse(content_type="application/pdf")
        filename = f"task_{task.id}_{datetime.now().strftime('%Y%m%d')}.pdf"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'

        pisa_status = pisa.CreatePDF(BytesIO(html.encode("utf-8")), dest=response)

        if pisa_status.err:
            return HttpResponse("PDF generation error", status=500)

        return response

    # Default: render HTML page with print button
    return render(request, "reports/task_completion_view.html", context)


@report_auth_required
def send_task_email(request, task_id):
    """დასრულებული ტასკის PDF გაგზავნა ემაილზე"""
    from django.core.mail import EmailMessage
    from django.conf import settings

    task = get_object_or_404(Task, id=task_id)
    equipment = task.schedule.equipment
    site = equipment.site
    company = site.company

    if request.user.role == "manager":
        user_tenant = get_user_tenant(request.user)
        if company.tenant != user_tenant:
            return HttpResponse("Unauthorized", status=403)
    elif request.user.role == "client" and request.user.company != company:
        return HttpResponse("Unauthorized", status=403)
    elif request.user.role == "inspector" and task.assigned_to != request.user:
        return HttpResponse("Unauthorized", status=403)

    # Only approved tasks
    if task.status != "approved":
        return HttpResponse(
            "მხოლოდ დამტკიცებული ტასკების გაგზავნა შეიძლება", status=400
        )

    recipient_email = request.POST.get("email") or request.GET.get("email")
    if not recipient_email:
        return HttpResponse("Email მისამართი აუცილებელია", status=400)

    if not PDF_AVAILABLE:
        return HttpResponse("PDF არ არის ხელმისაწვდომი", status=503)

    context = {
        "task": task,
        "equipment": equipment,
        "site": site,
        "company": company,
        "report": getattr(task, "report", None),
        "photos": (
            task.report.photos.all()
            if hasattr(task, "report") and task.report
            else []
        ),
        "generated_at": timezone.now(),
    }

    _ensure_georgian_pdf_font()
    html = render_to_string("reports/task_completion_pdf.html", context)
    pdf_buffer = BytesIO()
    pisa.CreatePDF(BytesIO(html.encode("utf-8")), dest=pdf_buffer)
    pdf_buffer.seek(0)

    subject = f"სერვისის აქტი - {equipment.name} - {task.schedule.task_name}"
    body = f"""
გამარჯობა,

გაგზავნილია სერვისის აქტი შემდეგი სამუშაოსთვის:

კომპანია: {company.name}
ობიექტი: {site.name}
დანადგარი: {equipment.name}
სამუშაო: {task.schedule.task_name}
თარიღი: {task.completed_date or timezone.now()}

დოკუმენტი თანდართულია PDF ფორმატში.

პატივისცემით,
IMPOSTI Maintenance System
"""

    try:
        email = EmailMessage(
            subject=subject,
            body=body,
            from_email=settings.DEFAULT_FROM_EMAIL,
            to=[recipient_email],
        )
        email.attach(
            f"service_act_{task.id}.pdf", pdf_buffer.getvalue(), "application/pdf"
        )
        email.send()

        return HttpResponse("ემაილი წარმატებით გაიგზავნა!", status=200)
    except Exception as e:
        return HttpResponse(f"შეცდომა ემაილის გაგზავნისას: {str(e)}", status=500)


@report_auth_required
def generate_excel_report(request):
    if request.user.role not in ["manager", "client"]:
        return HttpResponse("Unauthorized", status=403)

    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks Report"

    headers = [
        "ID",
        "კომპანია",
        "ობიექტი",
        "დანადგარი",
        "შემოწმება",
        "სიხშირე",
        "ინსპექტორი",
        "სტატუსი",
        "დაგეგმილი თარიღი",
        "დასრულების თარიღი",
        "კომენტარი",
    ]
    ws.append(headers)

    header_fill = PatternFill(
        start_color="0066CC", end_color="0066CC", fill_type="solid"
    )
    header_font = Font(bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    tasks = Task.objects.select_related(
        "schedule__equipment__site__company", "assigned_to", "report"
    )
    if request.user.role == "client":
        if not request.user.company:
            tasks = tasks.none()
        else:
            tasks = tasks.filter(
                schedule__equipment__site__company=request.user.company
            )
    else:
        user_tenant = get_user_tenant(request.user)
        tasks = tasks.filter(
            schedule__equipment__site__company__tenant=user_tenant
        )

    tasks = tasks.order_by("-scheduled_date")

    for task in tasks:
        scheduled_date_str = (
            task.scheduled_date.strftime("%d.%m.%Y %H:%M")
            if getattr(task, "scheduled_date", None)
            else ""
        )
        completed_date_str = (
            task.completed_date.strftime("%d.%m.%Y %H:%M")
            if getattr(task, "completed_date", None)
            else ""
        )
        comment_text = ""
        if hasattr(task, "report") and getattr(task, "report") is not None:
            comment_text = getattr(task.report, "comment", "")

        ws.append(
            [
                task.id,
                task.schedule.equipment.site.company.name,
                task.schedule.equipment.site.name,
                task.schedule.equipment.name,
                task.schedule.task_name,
                task.schedule.get_frequency_display(),
                task.assigned_to.username if task.assigned_to else "",
                task.get_status_display(),
                scheduled_date_str,
                completed_date_str,
                comment_text,
            ]
        )

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="tasks_report_{datetime.now().strftime("%Y%m%d")}.xlsx"'
    )
    wb.save(response)

    return response
